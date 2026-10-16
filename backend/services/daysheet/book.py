"""Builds the day sheet workbook (.xlsx) that is filled in on a phone with free Excel.

Tabs: Sales, Expenses, Summary (the close-the-day form), plus hidden Products (the
catalogue and stock when the sheet was made), Lists (expense categories) and a very
hidden Meta tab with the sheet's id and version. The free Excel app runs no macros,
so all the checking the sheet does on its own is plain formulas and colours; the POS
re-checks everything on upload.
"""
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from services.daysheet.layout import (
    CARD_CELL, COUNTED_CELL, EXP_AMOUNT, EXP_CATEGORY, EXP_DESCRIPTION, EXP_PAID_FROM, FIRST, LAST,
    LISTS_TAB, META_TAB, PRODUCTS_TAB, SALES_PRICE, SALES_PRODUCT, SALES_QTY, START_CELL,
    TEMPLATE_VERSION, CatalogItem, SheetInputs, SheetMeta,
)
from services.daysheet.strings import T

RED = PatternFill("solid", fgColor="F8B4B4")
AMBER = PatternFill("solid", fgColor="FDE2A7")
GREEN = PatternFill("solid", fgColor="C6EFCE")
GREY = PatternFill("solid", fgColor="E7E6E6")
HEAD = PatternFill("solid", fgColor="1F3864")
INPUT = PatternFill("solid", fgColor="FFF2CC")
WHITE_B = Font(bold=True, color="FFFFFF", size=12)
BIG = Font(bold=True, size=14)
LINK = Font(bold=True, color="FFFF00", size=12, underline="single")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
MONEY = "#,##0.00"


def build_workbook(
    meta: SheetMeta,
    day: date,
    issued_at: datetime,
    catalog: list[CatalogItem],
    categories: list[str],
    inputs: SheetInputs | None = None,
) -> bytes:
    t = T[meta.lang]
    wb = Workbook()
    names, count = _products(wb, catalog)
    sales = _sales(wb, t, names, count)
    expenses = _expenses(wb, t, categories)
    summary = _summary(wb, t, meta, day, issued_at)
    _meta(wb, meta, day)
    for name in (PRODUCTS_TAB, LISTS_TAB, META_TAB):
        wb.move_sheet(name, offset=10)
    if inputs:
        _fill(t, sales, expenses, summary, inputs)
    wb.active = 0
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def file_name(meta: SheetMeta, day: date) -> str:
    return f"DaySheet-{day:%m%d}-{meta.code.rsplit('-', 1)[-1]}-v{meta.version}.xlsx"


def _q(text: str) -> str:
    """Excel string literal."""
    return '"' + text.replace('"', '""') + '"'


def _esc(ref: str) -> str:
    """Escape Excel wildcards so names like '36PC*45GM' match literally."""
    return f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(TRIM({ref}),"~","~~"),"*","~*"),"?","~?")'


def _header(ws: Worksheet, titles: list[str], widths: list[int]) -> None:
    for col, (title, width) in enumerate(zip(titles, widths), start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill, cell.font = HEAD, WHITE_B
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"


def _link(cell, target: str, text: str) -> None:
    cell.value = text
    cell.hyperlink = f"#'{target}'!A1"
    cell.font = LINK
    cell.fill = HEAD


def _numbered(ws: Worksheet, r: int) -> None:
    ws[f"A{r}"] = r - 1
    ws[f"A{r}"].alignment = Alignment(horizontal="center")
    ws[f"A{r}"].font = Font(color="808080")
    for c in "ABCDEF":
        ws[f"{c}{r}"].border = BOX


def _products(wb: Workbook, catalog: list[CatalogItem]) -> tuple[str, int]:
    ws = wb.create_sheet(PRODUCTS_TAB)
    ws.append(["id", "name", "fixed_price", "stock"])
    for item in catalog:
        ws.append([item.id, item.name, float(item.fixed_price), float(item.stock)])
    ws.sheet_state = "hidden"
    ws.protection.sheet = True
    last = max(len(catalog), 1) + 1
    return f"{PRODUCTS_TAB}!$B$2:$B${last}", last


def _sales(wb: Workbook, t: dict, names: str, plast: int) -> Worksheet:
    """Hidden helper columns: H typed text (wildcards escaped), I matched position, J-M the
    product's id/name/fixed price/stock, N quantity of it sold so far, O how many names
    contain the text, P the price used."""
    ws = wb.active
    ws.title = t["sales"]
    _header(ws, t["sales_head"], [5, 30, 7, 9, 10, 30])
    _link(ws["F1"], t["summary"], t["to_summary"])
    col = lambda c: f"{PRODUCTS_TAB}!${c}$2:${c}${plast}"
    for r in range(FIRST, LAST + 1):
        _numbered(ws, r)
        ws[f"{SALES_PRICE}{r}"] = f'=IF(AND($J{r}<>"",N($L{r})>0),$L{r},"")'
        ws[f"E{r}"] = f'=IF(AND(J{r}<>"",ISNUMBER(C{r}),ISNUMBER(P{r})),C{r}*P{r},"")'
        ws[f"F{r}"] = (
            f'=IF(AND(TRIM(B{r})="",C{r}="",D{r}=""),"",'
            f'IF(TRIM(B{r})="",{_q(t["no_product"])},'
            f'IF(I{r}="",IF(N(O{r})>1,"⚠ "&O{r}&{_q(t["matches"])},{_q(t["not_found"])}),'
            f'IF(OR(NOT(ISNUMBER(C{r})),N(C{r})<=0),{_q(t["no_qty"])},'
            f'IF(N(P{r})<=0,{_q(t["no_price"])},'
            f'IF(N{r}>M{r},{_q(t["no_stock"])}&MAX(0,M{r}-(N{r}-C{r}))&{_q(t["left"])},"✓ "&K{r}))))))'
        )
        ws[f"H{r}"] = f"={_esc(f'B{r}')}"
        ws[f"I{r}"] = (
            f'=IF(TRIM(B{r})="","",IFERROR(MATCH(H{r},{names},0),'
            f'IF(COUNTIF({names},H{r}&"*")=1,MATCH(H{r}&"*",{names},0),'
            f'IF(COUNTIF({names},"*"&H{r}&"*")=1,MATCH("*"&H{r}&"*",{names},0),""))))'
        )
        ws[f"J{r}"] = f'=IF(I{r}="","",INDEX({col("A")},I{r}))'
        ws[f"K{r}"] = f'=IF(I{r}="","",INDEX({col("B")},I{r}))'
        ws[f"L{r}"] = f'=IF(I{r}="","",INDEX({col("C")},I{r}))'
        ws[f"M{r}"] = f'=IF(I{r}="","",INDEX({col("D")},I{r}))'
        ws[f"N{r}"] = f'=IF(J{r}="","",SUMIF($J${FIRST}:J{r},J{r},$C${FIRST}:C{r}))'
        ws[f"O{r}"] = f'=IF(OR(TRIM(B{r})="",I{r}<>""),"",COUNTIF({names},"*"&H{r}&"*"))'
        ws[f"P{r}"] = f'=IF(ISNUMBER(D{r}),D{r},IF(N(L{r})>0,L{r},""))'
        ws[f"D{r}"].number_format = ws[f"E{r}"].number_format = MONEY
    for c in "GHIJKLMNOP":
        ws.column_dimensions[c].hidden = True
    count = plast - 1
    # The list shrinks to the names starting with what was typed (after Enter).
    narrowing = (
        f'OFFSET({PRODUCTS_TAB}!$B$1,IFERROR(MATCH($H{FIRST}&"*",{names},0),1),0,'
        f'IF(COUNTIF({names},$H{FIRST}&"*"),COUNTIF({names},$H{FIRST}&"*"),{count}),1)'
    )
    dv = DataValidation(type="list", formula1=narrowing, allow_blank=True, showErrorMessage=False)
    dv.add(f"{SALES_PRODUCT}{FIRST}:{SALES_PRODUCT}{LAST}")
    ws.add_data_validation(dv)
    area = f"A{FIRST}:F{LAST}"
    ws.conditional_formatting.add(area, FormulaRule(formula=[f'LEFT($F{FIRST},1)="✗"'], fill=RED, font=Font(bold=True, color="9C0006")))
    ws.conditional_formatting.add(area, FormulaRule(formula=[f'LEFT($F{FIRST},1)="⚠"'], fill=AMBER))
    ws.conditional_formatting.add(f"F{FIRST}:F{LAST}", FormulaRule(formula=[f'LEFT($F{FIRST},1)="✓"'], font=Font(color="006100")))
    return ws


def _expenses(wb: Workbook, t: dict, categories: list[str]) -> Worksheet:
    """Hidden: H category position (0 = none), I paid-from in English, J/K drawer/other amount."""
    ws = wb.create_sheet(t["expenses"])
    _header(ws, t["exp_head"], [5, 16, 10, 26, 11, 24])
    _link(ws["F1"], t["summary"], t["to_summary"])
    lists = wb.create_sheet(LISTS_TAB)
    lists.append(["category"])
    for name in categories:
        lists.append([name])
    lists.sheet_state = "hidden"
    lists.protection.sheet = True
    cats = f"{LISTS_TAB}!$A$2:$A${max(len(categories), 1) + 1}"
    drawer, other = t["drawer"], t["other"]
    for r in range(FIRST, LAST + 1):
        _numbered(ws, r)
        ws[f"H{r}"] = f'=IF(TRIM(B{r})="",0,IFERROR(MATCH(TRIM(B{r}),{cats},0),""))'
        ws[f"I{r}"] = f'=IF(E{r}={_q(drawer)},"Drawer",IF(E{r}={_q(other)},"Other",""))'
        ws[f"F{r}"] = (
            f'=IF(AND(TRIM(B{r})="",C{r}="",TRIM(D{r})="",TRIM(E{r})=""),"",'
            f'IF(H{r}="",{_q(t["unknown_cat"])},'
            f'IF(OR(NOT(ISNUMBER(C{r})),N(C{r})<=0),{_q(t["no_amount"])},'
            f'IF(I{r}="",{_q(t["paid_from_q"])},"✓ "&E{r}&IF(H{r}=0,""," · "&INDEX({cats},H{r}))))))'
        )
        ws[f"J{r}"] = f'=IF(AND(LEFT(F{r},1)="✓",I{r}="Drawer"),C{r},0)'
        ws[f"K{r}"] = f'=IF(AND(LEFT(F{r},1)="✓",I{r}="Other"),C{r},0)'
        ws[f"C{r}"].number_format = MONEY
    for c in "GHIJK":
        ws.column_dimensions[c].hidden = True
    # Lists must be ranges or a quoted literal: the free phone app ignores other formula lists.
    for formula, column, title in (
        (cats, EXP_CATEGORY, t["pick_cat"]),
        (f'"{drawer},{other}"', EXP_PAID_FROM, t["pick_paid"]),
    ):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                            errorTitle=title, error=title)
        dv.add(f"{column}{FIRST}:{column}{LAST}")
        ws.add_data_validation(dv)
    ws.conditional_formatting.add(f"A{FIRST}:F{LAST}", FormulaRule(formula=[f'LEFT($F{FIRST},1)="⚠"'], fill=AMBER))
    ws.conditional_formatting.add(f"F{FIRST}:F{LAST}", FormulaRule(formula=[f'LEFT($F{FIRST},1)="✓"'], font=Font(color="006100")))
    return ws


def _summary(wb: Workbook, t: dict, meta: SheetMeta, day: date, issued_at: datetime) -> Worksheet:
    ws = wb.create_sheet(t["summary"])
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 22
    _link(ws["A1"], t["sales"], t["back_sales"])
    _link(ws["B1"], t["expenses"], t["back_exp"])
    ws.row_dimensions[1].height = 26
    sales, exp = f"'{t['sales']}'", f"'{t['expenses']}'"
    S, E = f"{sales}!$F${FIRST}:$F${LAST}", f"{exp}!$F${FIRST}:$F${LAST}"
    col = lambda c: f"{exp}!${c}${FIRST}:${c}${LAST}"
    issued = f"{meta.version} · {t['s_issued']} {issued_at.strftime(t['date_fmt'] + ' %H:%M')}"
    rows = {
        2: (t["s_sheet"], meta.code, "id"),
        3: (t["s_version"], issued, None),
        4: (t["s_date"], day.strftime(t["date_fmt"]), None),
        6: (t["s_sales"], None, "section"),
        7: (t["s_grand"], f'=SUMIF({S},"✓*",{sales}!$E${FIRST}:$E${LAST})', "money"),
        8: (t["s_card"], None, "input"),
        9: (t["s_cash"], f"=B7-N({CARD_CELL})", "total"),
        11: (t["s_check"], None, "section"),
        12: (t["s_counted"], None, "input"),
        13: (t["s_drawer"], f"=SUM({col('J')})", "money"),
        14: (t["s_start"], None, "input"),
        15: (t["s_accounted"], f"=N({COUNTED_CELL})+B13-N({START_CELL})", "total"),
        17: (t["s_diff"], f'=IF({COUNTED_CELL}="","",ROUND(B15-B9,2))', "diff"),
        18: ("", f'=IF(B17="",{_q(t["d_enter"])},IF(B17>0,{_q(t["d_extra"])},IF(B17<0,{_q(t["d_short"])},{_q(t["d_ok"])})))', "difflabel"),
        20: (t["s_other"], f"=SUM({col('K')})", "money"),
        21: (t["s_total_exp"], "=B13+B20", "money"),
        23: (t["s_ok"], f'=COUNTIF({S},"✓*")', None),
        24: (t["s_problems"], f'=COUNTIF({S},"⚠*")+COUNTIF({S},"✗*")+COUNTIF({E},"⚠*")', "problems"),
        25: (t["s_ready"], (
            f'=IF(B24>0,{_q(t["r_fix"])}&B24&{_q(t["r_rows"])},'
            f'IF(AND(B23>0,{CARD_CELL}=""),{_q(t["r_card"])},'
            f'IF(N({CARD_CELL})>B7,{_q(t["r_card_big"])},'
            f'IF(AND(B23>0,{COUNTED_CELL}=""),{_q(t["r_count"])},{_q(t["r_yes"])}))))'
        ), "ready"),
    }
    for r, (label, value, kind) in rows.items():
        a, b = ws[f"A{r}"], ws[f"B{r}"]
        a.value, b.value = label, value
        a.data_type = "s"
        a.font = b.font = Font(size=12)
        b.protection = Protection(locked=kind != "input")
        if kind == "section":
            a.font, a.fill, b.fill = WHITE_B, HEAD, HEAD
            continue
        a.border = b.border = BOX
        if kind in ("money", "total", "input", "diff"):
            b.number_format = MONEY
        if kind == "input":
            b.fill = INPUT
        if kind in ("total", "diff", "ready", "id"):
            a.font = b.font = BIG
        b.alignment = Alignment(horizontal="right")
    diff = "$B$17"
    ws.conditional_formatting.add("B17:B18", FormulaRule(formula=[f"AND(ISNUMBER({diff}),{diff}>0)"], fill=GREEN, font=Font(bold=True, color="006100")))
    ws.conditional_formatting.add("B17:B18", FormulaRule(formula=[f"AND(ISNUMBER({diff}),{diff}<0)"], fill=RED, font=Font(bold=True, color="9C0006")))
    ws.conditional_formatting.add("B17:B18", FormulaRule(formula=[f"AND(ISNUMBER({diff}),{diff}=0)"], fill=GREY))
    ws.conditional_formatting.add("B24", FormulaRule(formula=["$B$24>0"], fill=RED))
    ws.conditional_formatting.add("B25", FormulaRule(formula=['LEFT($B$25,1)="✓"'], fill=GREEN, font=Font(bold=True, color="006100")))
    ws.conditional_formatting.add("B25", FormulaRule(formula=['LEFT($B$25,1)="✗"'], fill=RED, font=Font(bold=True, color="9C0006")))
    ws.protection.sheet = True
    return ws


def _meta(wb: Workbook, meta: SheetMeta, day: date) -> None:
    ws = wb.create_sheet(META_TAB)
    for row in (
        ("sheet_id", meta.token),
        ("code", meta.code),
        ("version", meta.version),
        ("day", day.isoformat()),
        ("lang", meta.lang),
        ("template_version", TEMPLATE_VERSION),
        ("source", "KEYPOS"),
    ):
        ws.append(row)
    ws.sheet_state = "veryHidden"


def _fill(t: dict, sales: Worksheet, expenses: Worksheet, summary: Worksheet, inputs: SheetInputs) -> None:
    """Write back what was typed into an earlier version, in the same rows."""
    for line in inputs.sales:
        sales[f"{SALES_PRODUCT}{line.row}"] = line.product or None
        sales[f"{SALES_QTY}{line.row}"] = line.qty
        if line.price is not None:
            sales[f"{SALES_PRICE}{line.row}"] = line.price
    for line in inputs.expenses:
        paid = {"drawer": t["drawer"], "other": t["other"]}.get(line.paid_from, line.paid_from)
        expenses[f"{EXP_CATEGORY}{line.row}"] = line.category or None
        expenses[f"{EXP_AMOUNT}{line.row}"] = line.amount
        expenses[f"{EXP_DESCRIPTION}{line.row}"] = line.description or None
        expenses[f"{EXP_PAID_FROM}{line.row}"] = paid or None
    summary[CARD_CELL] = inputs.card
    summary[COUNTED_CELL] = inputs.counted
    summary[START_CELL] = inputs.start
