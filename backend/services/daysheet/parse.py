"""Reads a day sheet back: its id and version, the catalogue it was made with, and what was typed.

Only typed cells are read. Formula results (status, totals, matched product) are ignored
because the POS works them out again itself, so the upload doesn't depend on the phone
app having recalculated or saved them.
"""
import zipfile
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from services.daysheet.layout import (
    CARD_CELL, COUNTED_CELL, EXP_AMOUNT, EXP_CATEGORY, EXP_DESCRIPTION, EXP_PAID_FROM, FIRST, LAST,
    META_TAB, PRODUCTS_TAB, SALES_PRICE, SALES_PRODUCT, SALES_QTY, START_CELL,
    CatalogItem, ExpenseInput, Raw, SalesInput, SheetInputs, SheetMeta,
)
from services.daysheet.strings import T

MAX_UNPACKED = 60 * 1024 * 1024
PAID_FROM = {
    word.casefold(): key
    for strings in T.values()
    for key, word in (("drawer", strings["drawer"]), ("other", strings["other"]))
} | {"drawer": "drawer", "other": "other"}


class SheetError(ValueError):
    """The file isn't a readable KEYPOS day sheet."""


def read_sheet(data: bytes) -> tuple[SheetMeta, list[CatalogItem], SheetInputs]:
    wb = _open(data)
    if META_TAB not in wb.sheetnames:
        raise SheetError("This isn't a KEYPOS day sheet. Download one from the Day Sheet screen.")
    meta = _meta(wb[META_TAB])
    t = T[meta.lang]
    missing = [name for name in (t["sales"], t["expenses"], t["summary"], PRODUCTS_TAB) if name not in wb.sheetnames]
    if missing:
        raise SheetError(f"The sheet is missing its {missing[0]} tab. Download a new sheet.")
    summary = wb[t["summary"]]
    inputs = SheetInputs(
        sales=_sales(wb[t["sales"]]),
        expenses=_expenses(wb[t["expenses"]]),
        card=_typed(summary[CARD_CELL].value),
        counted=_typed(summary[COUNTED_CELL].value),
        start=_typed(summary[START_CELL].value),
    )
    return meta, _catalog(wb[PRODUCTS_TAB]), inputs


def to_decimal(raw: Raw) -> Decimal | None:
    """A typed number; "1,5" (comma decimal) is accepted. None when blank, ValueError when not a number."""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None
    if isinstance(raw, bool):
        raise ValueError(raw)
    text = str(raw).strip().replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        value = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(raw) from exc
    if not value.is_finite():
        raise ValueError(raw)
    return value


def paid_from_key(text: str) -> str | None:
    return PAID_FROM.get(text.strip().casefold()) if text else None


def _open(data: bytes):
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNPACKED:
                raise SheetError("The file is too large to be a day sheet.")
        return load_workbook(BytesIO(data), data_only=False)
    except SheetError:
        raise
    except Exception as exc:
        raise SheetError("Couldn't read this file. Upload the .xlsx day sheet you downloaded from the POS.") from exc


def _meta(ws: Worksheet) -> SheetMeta:
    values = {str(row[0]): row[1] for row in ws.iter_rows(values_only=True) if row and row[0] is not None}
    try:
        meta = SheetMeta(
            token=str(values["sheet_id"]),
            code=str(values["code"]),
            version=int(values["version"]),
            lang=str(values["lang"]),
        )
    except (KeyError, TypeError, ValueError) as exc:
        raise SheetError("This day sheet was made by an older test version. Download a new sheet.") from exc
    if meta.lang not in T:
        raise SheetError("This day sheet's language isn't supported. Download a new sheet.")
    return meta


def _catalog(ws: Worksheet) -> list[CatalogItem]:
    items = []
    for row in ws.iter_rows(min_row=2, max_col=4, values_only=True):
        if not row or row[0] is None or row[1] is None:
            continue
        try:
            items.append(CatalogItem(
                id=int(row[0]), name=str(row[1]),
                fixed_price=to_decimal(row[2]) or Decimal("0"), stock=to_decimal(row[3]) or Decimal("0"),
            ))
        except ValueError:
            continue
    return items


def _typed(value: Raw) -> Raw:
    """None for formulas (cells the person didn't type into) and blanks."""
    if isinstance(value, str):
        if value.startswith("="):
            return None
        return value if value.strip() else None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    return None if value is None else str(value)


def _text(value: Raw) -> str:
    typed = _typed(value)
    if typed is None:
        return ""
    if isinstance(typed, float) and typed.is_integer():
        return str(int(typed))
    return str(typed).strip()


def _sales(ws: Worksheet) -> list[SalesInput]:
    rows = []
    for r in range(FIRST, LAST + 1):
        product = _text(ws[f"{SALES_PRODUCT}{r}"].value)
        qty = _typed(ws[f"{SALES_QTY}{r}"].value)
        price = _typed(ws[f"{SALES_PRICE}{r}"].value)
        if product or qty is not None or price is not None:
            rows.append(SalesInput(row=r, product=product, qty=qty, price=price))
    return rows


def _expenses(ws: Worksheet) -> list[ExpenseInput]:
    rows = []
    for r in range(FIRST, LAST + 1):
        category = _text(ws[f"{EXP_CATEGORY}{r}"].value)
        amount = _typed(ws[f"{EXP_AMOUNT}{r}"].value)
        description = _text(ws[f"{EXP_DESCRIPTION}{r}"].value)
        paid_raw = _text(ws[f"{EXP_PAID_FROM}{r}"].value)
        if category or amount is not None or description or paid_raw:
            rows.append(ExpenseInput(
                row=r, category=category, amount=amount, description=description,
                paid_from=paid_from_key(paid_raw) or paid_raw,
            ))
    return rows
