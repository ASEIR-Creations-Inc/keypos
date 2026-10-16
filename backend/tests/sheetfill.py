"""Fill a day sheet's typed cells the way a person does on a phone (used by pytest and the browser test)."""
import json
import sys
from io import BytesIO

from openpyxl import load_workbook

TABS = {"en": ("Sales", "Expenses", "Summary"), "pt": ("Vendas", "Despesas", "Resumo")}


def fill(data: bytes, sales=(), expenses=(), card=None, counted=None, start=None) -> bytes:
    """sales: (row, product, qty, price|None); expenses: (row, category, amount, description, paid_from)."""
    wb = load_workbook(BytesIO(data))
    lang = {r[0]: r[1] for r in wb["Meta"].iter_rows(values_only=True)}["lang"]
    ws_sales, ws_exp, ws_sum = (wb[name] for name in TABS[lang])
    for row, product, qty, price in sales:
        ws_sales[f"B{row}"], ws_sales[f"C{row}"] = product, qty
        if price is not None:
            ws_sales[f"D{row}"] = price
    for row, category, amount, description, paid in expenses:
        for col, value in zip("BCDE", (category, amount, description, paid)):
            ws_exp[f"{col}{row}"] = value
    for cell, value in (("B8", card), ("B12", counted), ("B14", start)):
        if value is not None:
            ws_sum[cell] = value
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


if __name__ == "__main__":
    # sheetfill.py IN OUT '{"sales": [[row, product, qty, price]], "expenses": [...], "card": 0, "counted": 0}'
    source, target, spec = sys.argv[1], sys.argv[2], json.loads(sys.argv[3])
    with open(source, "rb") as f:
        filled = fill(f.read(), spec.get("sales", ()), spec.get("expenses", ()), spec.get("card"), spec.get("counted"), spec.get("start"))
    with open(target, "wb") as f:
        f.write(filled)
