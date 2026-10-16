"""Day sheets: issue, fill (as a phone would), inspect, refresh, upload."""
import os
from decimal import Decimal
from io import BytesIO

import httpx
from openpyxl import load_workbook

from tests.conftest import stock_of
from tests.sheetfill import fill

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def new_sheet(client: httpx.AsyncClient, lang: str = "en") -> tuple[bytes, str]:
    response = await client.post(f"/daysheets?lang={lang}")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    return response.content, response.headers["x-day-sheet"]


def upload_file(data: bytes, name: str = "DaySheet.xlsx") -> dict:
    return {"file": (name, data, XLSX)}


async def inspect(client: httpx.AsyncClient, data: bytes) -> dict:
    response = await client.post("/daysheets/inspect", files=upload_file(data))
    assert response.status_code == 200, response.text
    return response.json()


async def category(client: httpx.AsyncClient, name: str) -> dict:
    existing = {c["name"]: c for c in (await client.get("/expenses/categories")).json()}
    if name in existing:
        return existing[name]
    response = await client.post("/expenses/categories", json={"name": name})
    assert response.status_code == 201, response.text
    return response.json()


def messages(result: dict) -> list[str]:
    return [p["message"] for p in result["problems"]]


async def test_issue_builds_a_workbook_tied_to_the_pos(client, make_product):
    product = await make_product(stock="7")
    data, code = await new_sheet(client, "pt")
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames[:3] == ["Vendas", "Despesas", "Resumo"]
    meta = {r[0]: r[1] for r in wb["Meta"].iter_rows(values_only=True)}
    assert meta["code"] == code and meta["version"] == 1 and meta["lang"] == "pt"
    assert wb["Resumo"]["B2"].value == code
    names = [r[1] for r in wb["Products"].iter_rows(min_row=2, values_only=True)]
    assert product["name"] in names

    listed = (await client.get("/daysheets")).json()
    assert listed[0]["code"] == code and listed[0]["status"] == "open"


async def test_blank_sheet_cannot_be_uploaded(client, shift):
    data, _ = await new_sheet(client)
    result = await inspect(client, data)
    assert not result["can_upload"] and result["can_refresh"]
    assert "The sheet has no sales or expenses yet" in messages(result)


async def test_upload_records_one_receipt_and_the_expenses(client, shift, make_product):
    a = await make_product(stock="10", cost="1.00")
    b = await make_product(stock="5", cost="2.00")
    daily = await category(client, "Daily")
    data, code = await new_sheet(client)
    data = fill(
        data,
        sales=[(2, a["name"], 2, 1.5), (3, b["name"][:-1], 1, "12,50"), (5, a["name"].upper(), 3, 1.5)],
        expenses=[(2, "daily", 15, "Ice", "Drawer"), (3, None, 300, "Rent", "Other")],
        card=10, counted=20.5, start=5,
    )
    result = await inspect(client, data)
    assert result["problems"] == [] and result["can_upload"], result["problems"]
    totals = {k: Decimal(str(v)) if v is not None else None for k, v in result["totals"].items()}
    assert totals["grand_total"] == Decimal("20.00")
    assert totals["cash_sales"] == Decimal("10.00")
    assert totals["drawer_expenses"] == Decimal("15.00") and totals["other_expenses"] == Decimal("300.00")
    assert totals["accounted"] == Decimal("30.50")  # 20.50 counted + 15 drawer - 5 start
    assert totals["difference"] == Decimal("20.50")
    assert [l["product_id"] for l in result["lines"]] == [a["id"], b["id"], a["id"]]

    response = await client.post("/daysheets/upload", files=upload_file(data, "DaySheet-test.xlsx"))
    assert response.status_code == 200, response.text
    receipt = response.json()["receipt_ref"]
    assert receipt.startswith("TXN-") and response.json()["expense_count"] == 2

    assert await stock_of(client, a["id"]) == Decimal("5")
    assert await stock_of(client, b["id"]) == Decimal("4")
    txns = (await client.get("/transactions", params={"limit": 5})).json()
    txn = next(t for t in (txns["items"] if isinstance(txns, dict) else txns) if t["ref_number"] == receipt)
    assert Decimal(str(txn["payment_cash"])) == Decimal("10.00")
    assert Decimal(str(txn["payment_card"])) == Decimal("10.00")
    assert code in txn["notes"]

    expenses = (await client.get("/expenses")).json()["items"]
    from_sheet = [e for e in expenses if e["day_sheet_id"]]
    assert {(e["description"], e["paid_from_drawer"], e["category_name"]) for e in from_sheet} >= {
        ("Ice", True, daily["name"]), ("Rent", False, None),
    }

    trail = (await client.get(f"/shifts/{shift['id']}/report")).json()
    sheet_entries = [e for e in trail["activity"] if e["action"] == "day_sheet_uploaded"]
    assert sheet_entries and "+20.50" in sheet_entries[-1]["summary"]

    again = await client.post("/daysheets/upload", files=upload_file(data))
    assert again.status_code == 400 and "already uploaded" in again.json()["detail"]
    listed = next(s for s in (await client.get("/daysheets")).json() if s["code"] == code)
    assert listed["status"] == "uploaded" and listed["receipt_ref"] == receipt


async def test_stock_shortfall_blocks_everything(client, shift, make_product):
    a = await make_product(stock="3")
    data, _ = await new_sheet(client)
    data = fill(data, sales=[(2, a["name"], 2, 1), (3, a["name"], 2, 1)],
                expenses=[(2, None, 4, "Bread", "Drawer")], card=0, counted=0)
    result = await inspect(client, data)
    assert not result["can_upload"]
    assert any("the sheet sells 4, stock is 3" in m for m in messages(result)), messages(result)
    response = await client.post("/daysheets/upload", files=upload_file(data))
    assert response.status_code == 400 and "stock is 3" in response.json()["detail"]
    assert await stock_of(client, a["id"]) == Decimal("3")
    expenses = (await client.get("/expenses")).json()["items"]
    assert not any(e["description"] == "Bread" for e in expenses)


async def test_row_problems_are_reported_per_line(client, shift, make_product):
    tag = os.urandom(3).hex()
    await make_product(name=f"Twin {tag} one")
    await make_product(name=f"Twin {tag} two")
    data, _ = await new_sheet(client)
    data = fill(
        data,
        sales=[(2, f"Twin {tag}", 1, 1), (3, "No such product xyz", 1, 1), (4, f"Twin {tag} one", None, 1),
               (5, f"Twin {tag} two", 1, None), (6, f"Twin {tag} two", "abc", 1)],
        expenses=[(2, "Transport", 5, None, "Drawer"), (3, None, None, "x", "Drawer"), (4, None, 2, None, "Bank")],
        card=50,
    )
    found = {(p["tab"], p["row"], p["message"]) for p in (await inspect(client, data))["problems"]}
    assert ("sales", 2, f'"Twin {tag}" matches 2 products: type more of the name') in found
    assert ("sales", 3, '"No such product xyz" was not found in the POS') in found
    assert ("sales", 4, "No quantity") in found
    assert ("sales", 5, "No price") in found
    assert ("sales", 6, 'Quantity "abc" isn\'t a number') in found
    assert ("expenses", 2, 'Category "Transport" doesn\'t exist in the POS') in found
    assert ("expenses", 3, "No amount") in found
    assert ("expenses", 4, "Paid from must be Drawer or Other") in found
    assert ("summary", None, "Enter the cash counted on the Summary tab") in found
    assert ("summary", None, "The card slip total 50.00 is more than the day's sales 0.00") in found


async def test_refresh_keeps_rows_and_retires_the_old_version(client, shift, make_product):
    a = await make_product(stock="2")
    data, code = await new_sheet(client, "pt")
    v1 = fill(data, sales=[(2, a["name"], 2, 3), (4, "typo here", 1, 1)],
              expenses=[(3, None, 7, "Gelo", "Caixa")], card=1, counted=5)
    late = await make_product(stock="9")
    await client.put(f"/products/{a['id']}", json={"name": a["name"] + " renamed"})

    response = await client.post("/daysheets/refresh", files=upload_file(v1))
    assert response.status_code == 200, response.text
    assert response.headers["x-day-sheet"] == code and response.headers["x-day-sheet-version"] == "2"
    assert "-v2.xlsx" in response.headers["content-disposition"]
    v2 = response.content
    wb = load_workbook(BytesIO(v2))
    sales, exp, summary = wb["Vendas"], wb["Despesas"], wb["Resumo"]
    assert sales["B2"].value == a["name"] + " renamed" and sales["C2"].value == 2 and sales["D2"].value == 3
    assert sales["B4"].value == "typo here"
    assert str(sales["D3"].value).startswith("=")
    assert exp["C3"].value == 7 and exp["D3"].value == "Gelo" and exp["E3"].value == "Caixa"
    assert summary["B8"].value == 1 and summary["B12"].value == 5
    assert late["name"] in [r[1] for r in wb["Products"].iter_rows(min_row=2, values_only=True)]

    stale = await inspect(client, v1)
    assert not stale["can_refresh"] and not stale["can_upload"]
    assert f"This is version 1 of {code}, but version 2 exists. Use the newest file." in messages(stale)
    assert (await client.post("/daysheets/refresh", files=upload_file(v1))).status_code == 400

    sheet_id = next(s["id"] for s in (await client.get("/daysheets")).json() if s["code"] == code)
    again = await client.get(f"/daysheets/{sheet_id}/file")
    assert again.status_code == 200 and again.content == v2


async def test_products_created_after_download_are_found_live(client, shift, make_product):
    data, _ = await new_sheet(client)
    late = await make_product(stock="4")
    data = fill(data, sales=[(2, late["name"], 1, 2)], card=0, counted=2)
    result = await inspect(client, data)
    assert result["can_upload"], result["problems"]
    assert result["lines"][0]["product_id"] == late["id"]


async def test_expenses_only_sheet(client, shift):
    data, _ = await new_sheet(client)
    data = fill(data, expenses=[(2, None, 3, "Water", "Drawer")])
    result = await inspect(client, data)
    assert result["can_upload"], result["problems"]
    response = await client.post("/daysheets/upload", files=upload_file(data))
    assert response.status_code == 200 and response.json()["receipt_ref"] is None


async def test_upload_needs_an_open_shift(client, make_product):
    await client.post("/shifts/close", json={"closing_cash": 0})
    a = await make_product()
    data, _ = await new_sheet(client)
    data = fill(data, sales=[(2, a["name"], 1, 1)], card=0, counted=1)
    assert "Open a shift before uploading" in messages(await inspect(client, data))
    assert (await client.post("/daysheets/upload", files=upload_file(data))).status_code == 400


async def test_sheet_expenses_are_admin_only(client, staff, make_product):
    cashier = await staff("cashier")
    await cashier.post("/shifts/open", json={"opening_float": 0})
    data, _ = await new_sheet(cashier)
    data = fill(data, expenses=[(2, None, 6, "Fuel", "Drawer")])
    response = await cashier.post("/daysheets/upload", files=upload_file(data))
    assert response.status_code == 200, response.text
    expense = next(e for e in (await cashier.get("/expenses")).json()["items"] if e["description"] == "Fuel")
    assert expense["day_sheet_id"]
    denied = await cashier.delete(f"/expenses/{expense['id']}")
    assert denied.status_code == 403 and "only an admin" in denied.json()["detail"]
    assert (await client.delete(f"/expenses/{expense['id']}")).status_code == 200


async def test_foreign_files_are_refused(client):
    response = await client.post("/daysheets/inspect", files=upload_file(b"not a spreadsheet"))
    assert response.status_code == 400 and "Couldn't read this file" in response.json()["detail"]
    data, _ = await new_sheet(client)
    wb = load_workbook(BytesIO(data))
    wb["Meta"]["B1"] = "00000000-0000-0000-0000-000000000000"
    out = BytesIO()
    wb.save(out)
    response = await client.post("/daysheets/inspect", files=upload_file(out.getvalue()))
    assert response.status_code == 400 and "wasn't issued by this POS" in response.json()["detail"]
